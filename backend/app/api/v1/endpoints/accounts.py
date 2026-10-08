import csv
import io
import re
from typing import Optional, List
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Response, status
from sqlalchemy import select, func, or_
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.models.company import Company
from app.models.account import Account
from app.models.supplier import Supplier
from app.models.invoice import Invoice
from app.models.accounting_entry import AccountingEntryLine
from app.schemas.account_dto import (
    AccountCreateDTO,
    AccountUpdateDTO,
    AccountResponseDTO,
    AccountPaginatedResponseDTO,
    AccountImportSummaryDTO,
)
from app.services.nif_validator import normalize_nif

router = APIRouter()

def deduce_account_type(codigo: str) -> str:
    """Deduce el tipo de cuenta según el cuadro de cuentas del PGC español."""
    c = codigo.strip()
    if c.startswith("6"):
        return "GASTO"
    elif c.startswith("7"):
        return "INGRESO"
    elif c.startswith("400") or c.startswith("401"):
        return "PROVEEDOR"
    elif c.startswith("410") or c.startswith("411"):
        return "ACREEDOR"
    elif c.startswith("430") or c.startswith("440"):
        return "CLIENTE"
    elif c.startswith("57") or c.startswith("52") or c.startswith("17"):
        return "FINANCIERO"
    elif c.startswith("47"):
        return "TRIBUTARIO"
    return "OTRO"

def parse_float_val(val: Optional[object]) -> float:
    """Convierte cadenas con símbolos de moneda y comas a flotante seguro."""
    if val is None:
        return 0.0
    s = str(val).replace("€", "").replace("EUR", "").replace(" ", "").strip()
    if not s:
        return 0.0
    # Si contiene tanto punto como coma (ej. 1.250,50)
    if "." in s and "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return 0.0

@router.get("/{company_id}/accounts", response_model=AccountPaginatedResponseDTO)
async def list_company_accounts(
    company_id: str,
    search: Optional[str] = Query(None, description="Búsqueda por código, descripción o CIF"),
    tipo: Optional[str] = Query(None, description="Filtrar por tipo (GASTO, INGRESO, PROVEEDOR, etc.)"),
    sort_by: Optional[str] = Query("codigo", description="Ordenar por: codigo, descripcion, saldo, debe, haber"),
    sort_order: Optional[str] = Query("asc", description="Dirección: asc o desc"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    db: AsyncSession = Depends(get_db)
):
    """
    Lista las cuentas del Plan Contable PYME con saldos contables en tiempo real:
    - Debe Acumulado (€)
    - Haber Acumulado (€)
    - Saldo Actual (€) y naturaleza (DEUDOR, ACREEDOR, CERO)
    Soporta ordenación por saldo, búsqueda y filtrado.
    """
    base_query = select(Account).where(Account.company_id == company_id)

    if search:
        s = f"%{search.strip()}%"
        base_query = base_query.where(
            or_(
                Account.codigo.ilike(s),
                Account.descripcion.ilike(s),
                Account.cif_asociado.ilike(s)
            )
        )
    if tipo and tipo != "ALL":
        t = tipo.upper().strip()
        if "_" in t:
            t = t.split("_")[0]
        base_query = base_query.where(Account.tipo == t)

    # Contar total
    count_query = select(func.count()).select_from(base_query.subquery())
    total_res = await db.execute(count_query)
    total = total_res.scalar_one()

    # Si no hay cuentas para esta empresa y no hay filtros aplicados, sembrar el PGC PYMES automáticamente
    if total == 0 and not search and (not tipo or tipo == "ALL"):
        c_res = await db.execute(select(Company).where(Company.id == company_id))
        company = c_res.scalars().first()
        if company:
            from app.services.pyme_pgc_seed import seed_company_chart_of_accounts
            await seed_company_chart_of_accounts(db, company_id, company.plan_cuentas_longitud)
            # Recontar tras sembrar
            count_query = select(func.count()).select_from(base_query.subquery())
            total_res = await db.execute(count_query)
            total = total_res.scalar_one()

    # Sumas acumuladas de apuntes contables (AccountingEntryLine) para las facturas de esta empresa
    entries_stmt = (
        select(
            AccountingEntryLine.subcuenta,
            func.coalesce(func.sum(AccountingEntryLine.debe), 0.0).label("sum_debe"),
            func.coalesce(func.sum(AccountingEntryLine.haber), 0.0).label("sum_haber")
        )
        .join(Invoice, Invoice.id == AccountingEntryLine.invoice_id)
        .where(Invoice.company_id == company_id)
        .group_by(AccountingEntryLine.subcuenta)
    )
    entries_res = await db.execute(entries_stmt)
    entries_map = {row.subcuenta: (float(row.sum_debe), float(row.sum_haber)) for row in entries_res.all()}

    # Traer todas las cuentas filtradas para calcular balances y ordenar con precisión
    all_filtered_res = await db.execute(base_query)
    accounts = all_filtered_res.scalars().all()

    account_dtos: List[AccountResponseDTO] = []
    for a in accounts:
        e_debe, e_haber = entries_map.get(a.codigo, (0.0, 0.0))
        tot_debe = round(float(getattr(a, "debe_inicial", 0.0) or 0.0) + e_debe, 2)
        tot_haber = round(float(getattr(a, "haber_inicial", 0.0) or 0.0) + e_haber, 2)
        diff = round(tot_debe - tot_haber, 2)

        if diff > 0.005:
            saldo = diff
            tipo_saldo = "DEUDOR"
        elif diff < -0.005:
            saldo = round(abs(diff), 2)
            tipo_saldo = "ACREEDOR"
        else:
            saldo = 0.0
            tipo_saldo = "CERO"

        dto = AccountResponseDTO(
            id=a.id,
            company_id=a.company_id,
            codigo=a.codigo,
            descripcion=a.descripcion,
            tipo=a.tipo,
            cif_asociado=a.cif_asociado,
            debe_acumulado=tot_debe,
            haber_acumulado=tot_haber,
            saldo_actual=saldo,
            tipo_saldo=tipo_saldo,
            created_at=a.created_at
        )
        account_dtos.append(dto)

    # Ordenar
    is_desc = (sort_order or "asc").lower() == "desc"
    sb = (sort_by or "codigo").lower()
    if sb == "saldo":
        account_dtos.sort(key=lambda x: x.saldo_actual, reverse=is_desc)
    elif sb == "debe":
        account_dtos.sort(key=lambda x: x.debe_acumulado, reverse=is_desc)
    elif sb == "haber":
        account_dtos.sort(key=lambda x: x.haber_acumulado, reverse=is_desc)
    elif sb == "descripcion":
        account_dtos.sort(key=lambda x: x.descripcion.lower(), reverse=is_desc)
    else:  # codigo por defecto
        account_dtos.sort(key=lambda x: x.codigo, reverse=is_desc)

    # Paginación
    offset = (page - 1) * page_size
    paged_items = account_dtos[offset:offset + page_size]

    return AccountPaginatedResponseDTO(
        total=len(account_dtos),
        page=page,
        page_size=page_size,
        items=paged_items
    )

@router.post("/{company_id}/chart-of-accounts/seed")
async def seed_chart_of_accounts(
    company_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Precarga o complementa el catálogo completo del PGC PYMES para la empresa según sus dígitos."""
    c_res = await db.execute(select(Company).where(Company.id == company_id))
    company = c_res.scalars().first()
    if not company:
        raise HTTPException(status_code=404, detail="Empresa no encontrada.")

    from app.services.pyme_pgc_seed import seed_company_chart_of_accounts
    created_count = await seed_company_chart_of_accounts(db, company_id, company.plan_cuentas_longitud)

    return {
        "success": True,
        "message": f"Se han sembrado {created_count} subcuentas del PGC PYMES adaptadas a {company.plan_cuentas_longitud} dígitos.",
        "created_count": created_count,
        "plan_longitud": company.plan_cuentas_longitud
    }

@router.post("/{company_id}/accounts", response_model=AccountResponseDTO, status_code=status.HTTP_201_CREATED)
async def create_company_account(
    company_id: str,
    payload: AccountCreateDTO,
    db: AsyncSession = Depends(get_db)
):
    """Crea una subcuenta validando estrictamente la longitud configurada para la empresa."""
    c_res = await db.execute(select(Company).where(Company.id == company_id))
    company = c_res.scalars().first()
    if not company:
        raise HTTPException(status_code=404, detail="Empresa no encontrada.")

    clean_code = payload.codigo.strip()
    expected_len = company.plan_cuentas_longitud

    # Validación de longitud estricta
    if len(clean_code) != expected_len:
        raise HTTPException(
            status_code=400,
            detail=f"La subcuenta '{clean_code}' tiene {len(clean_code)} dígitos. El plan contable de esta empresa requiere exactamente {expected_len} dígitos."
        )

    # Comprobar duplicado
    dup_res = await db.execute(
        select(Account).where(Account.company_id == company_id, Account.codigo == clean_code)
    )
    if dup_res.scalars().first():
        raise HTTPException(status_code=400, detail=f"Ya existe una cuenta registrada con el código '{clean_code}'.")

    tipo_cuenta = payload.tipo or deduce_account_type(clean_code)
    cif_norm = normalize_nif(payload.cif_asociado) if payload.cif_asociado else None

    debe_ini = float(payload.debe_inicial or 0.0)
    haber_ini = float(payload.haber_inicial or 0.0)

    account = Account(
        company_id=company_id,
        codigo=clean_code,
        descripcion=payload.descripcion.strip(),
        tipo=tipo_cuenta,
        cif_asociado=cif_norm,
        debe_inicial=debe_ini,
        haber_inicial=haber_ini,
    )
    db.add(account)

    # Si es proveedor o acreedor con CIF, sincronizar con la tabla de proveedores
    if tipo_cuenta in ("PROVEEDOR", "ACREEDOR") and cif_norm:
        supp_res = await db.execute(
            select(Supplier).where(Supplier.company_id == company_id, Supplier.cif == cif_norm)
        )
        supplier = supp_res.scalars().first()
        if not supplier:
            new_supp = Supplier(
                company_id=company_id,
                cif=cif_norm,
                nombre=payload.descripcion.strip(),
                subcuenta_proveedor=clean_code,
                subcuenta_gasto_defecto="629" + "0" * (expected_len - 4) + "1"
            )
            db.add(new_supp)
        else:
            supplier.subcuenta_proveedor = clean_code

    await db.commit()
    await db.refresh(account)

    diff = round(debe_ini - haber_ini, 2)
    saldo = abs(diff)
    tipo_saldo = "DEUDOR" if diff > 0 else ("ACREEDOR" if diff < 0 else "CERO")

    return AccountResponseDTO(
        id=account.id,
        company_id=account.company_id,
        codigo=account.codigo,
        descripcion=account.descripcion,
        tipo=account.tipo,
        cif_asociado=account.cif_asociado,
        debe_acumulado=debe_ini,
        haber_acumulado=haber_ini,
        saldo_actual=saldo,
        tipo_saldo=tipo_saldo,
        created_at=account.created_at
    )

@router.put("/{company_id}/accounts/{account_id}", response_model=AccountResponseDTO)
async def update_company_account(
    company_id: str,
    account_id: str,
    payload: AccountUpdateDTO,
    db: AsyncSession = Depends(get_db)
):
    """Actualiza la descripción, tipo, CIF o saldos iniciales de una subcuenta."""
    res = await db.execute(
        select(Account).where(Account.id == account_id, Account.company_id == company_id)
    )
    account = res.scalars().first()
    if not account:
        raise HTTPException(status_code=404, detail="Cuenta no encontrada.")

    if payload.descripcion is not None:
        account.descripcion = payload.descripcion.strip()
    if payload.tipo is not None:
        account.tipo = payload.tipo
    if payload.cif_asociado is not None:
        account.cif_asociado = normalize_nif(payload.cif_asociado) if payload.cif_asociado else None
    if payload.debe_inicial is not None:
        account.debe_inicial = float(payload.debe_inicial)
    if payload.haber_inicial is not None:
        account.haber_inicial = float(payload.haber_inicial)

    await db.commit()
    await db.refresh(account)

    # Calcular saldos acumulados sumando apuntes contables
    entries_stmt = (
        select(
            func.coalesce(func.sum(AccountingEntryLine.debe), 0.0).label("sum_debe"),
            func.coalesce(func.sum(AccountingEntryLine.haber), 0.0).label("sum_haber")
        )
        .join(Invoice, Invoice.id == AccountingEntryLine.invoice_id)
        .where(Invoice.company_id == company_id, AccountingEntryLine.subcuenta == account.codigo)
    )
    entries_res = await db.execute(entries_stmt)
    row = entries_res.first()
    e_debe = float(row.sum_debe) if row else 0.0
    e_haber = float(row.sum_haber) if row else 0.0

    tot_debe = round(float(account.debe_inicial or 0.0) + e_debe, 2)
    tot_haber = round(float(account.haber_inicial or 0.0) + e_haber, 2)
    diff = round(tot_debe - tot_haber, 2)

    return AccountResponseDTO(
        id=account.id,
        company_id=account.company_id,
        codigo=account.codigo,
        descripcion=account.descripcion,
        tipo=account.tipo,
        cif_asociado=account.cif_asociado,
        debe_acumulado=tot_debe,
        haber_acumulado=tot_haber,
        saldo_actual=abs(diff),
        tipo_saldo="DEUDOR" if diff > 0 else ("ACREEDOR" if diff < 0 else "CERO"),
        created_at=account.created_at
    )

@router.delete("/{company_id}/accounts/{account_id}")
async def delete_company_account(
    company_id: str,
    account_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Elimina una cuenta del catálogo contable."""
    res = await db.execute(
        select(Account).where(Account.id == account_id, Account.company_id == company_id)
    )
    account = res.scalars().first()
    if not account:
        raise HTTPException(status_code=404, detail="Cuenta no encontrada.")

    await db.delete(account)
    await db.commit()
    return {"success": True, "message": "Cuenta eliminada correctamente."}

@router.post("/{company_id}/chart-of-accounts/import", response_model=AccountImportSummaryDTO)
async def import_chart_of_accounts(
    company_id: str,
    file: UploadFile = File(..., description="Fichero CSV o Excel (.xlsx/.xls) exportado de Contasol, A3 o Sage"),
    db: AsyncSession = Depends(get_db)
):
    """
    Importador masivo del Plan Contable PYME:
    Parsea columnas clave (código, descripción, CIF) y columnas opcionales de saldos (Debe, Haber, Saldo).
    Actualiza masivamente el catálogo y el maestro de proveedores.
    """
    c_res = await db.execute(select(Company).where(Company.id == company_id))
    company = c_res.scalars().first()
    if not company:
        raise HTTPException(status_code=404, detail="Empresa no encontrada.")

    expected_len = company.plan_cuentas_longitud
    contents = await file.read()
    filename = (file.filename or "").lower()

    rows: List[dict] = []
    errors: List[str] = []

    # 1. Parsear según formato (Excel o CSV)
    if filename.endswith(".xlsx") or filename.endswith(".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
            sheet = wb.active
            iter_rows = list(sheet.iter_rows(values_only=True))
            if not iter_rows:
                raise HTTPException(status_code=400, detail="El archivo Excel está vacío.")

            header = [str(col).strip().lower() if col is not None else "" for col in iter_rows[0]]
            code_idx, desc_idx, cif_idx = -1, -1, -1
            debe_idx, haber_idx, saldo_idx = -1, -1, -1

            for idx, h in enumerate(header):
                if any(k in h for k in ("cuenta", "codigo", "código", "subcuenta", "code")):
                    code_idx = idx
                elif any(k in h for k in ("descrip", "nombre", "concepto", "titulo", "name")):
                    desc_idx = idx
                elif any(k in h for k in ("cif", "nif", "dni", "fiscal")):
                    cif_idx = idx
                elif any(k in h for k in ("debe", "debit", "debe_acumulado", "debe_inicial")):
                    debe_idx = idx
                elif any(k in h for k in ("haber", "credit", "haber_acumulado", "haber_inicial")):
                    haber_idx = idx
                elif any(k in h for k in ("saldo", "balance")):
                    saldo_idx = idx

            if code_idx == -1:
                code_idx = 0
            if desc_idx == -1:
                desc_idx = 1 if len(header) > 1 else 0

            for r in iter_rows[1:]:
                if not r or r[code_idx] is None:
                    continue
                codigo_raw = str(r[code_idx]).strip()
                desc_raw = str(r[desc_idx]).strip() if desc_idx < len(r) and r[desc_idx] is not None else "Sin descripción"
                cif_raw = str(r[cif_idx]).strip() if cif_idx != -1 and cif_idx < len(r) and r[cif_idx] is not None else ""

                d_val = parse_float_val(r[debe_idx]) if debe_idx != -1 and debe_idx < len(r) else 0.0
                h_val = parse_float_val(r[haber_idx]) if haber_idx != -1 and haber_idx < len(r) else 0.0
                if d_val == 0.0 and h_val == 0.0 and saldo_idx != -1 and saldo_idx < len(r):
                    s_val = parse_float_val(r[saldo_idx])
                    if s_val > 0:
                        d_val = s_val
                    elif s_val < 0:
                        h_val = abs(s_val)

                rows.append({
                    "codigo": codigo_raw,
                    "descripcion": desc_raw,
                    "cif": cif_raw,
                    "debe": d_val,
                    "haber": h_val,
                })
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error leyendo archivo Excel: {str(e)}")

    else:
        # Modo CSV / Texto delimitado
        try:
            text = contents.decode("utf-8-sig", errors="replace")
            sample = text[:2000]
            delimiter = ";" if sample.count(";") >= sample.count(",") else ","
            if sample.count("\t") > max(sample.count(";"), sample.count(",")):
                delimiter = "\t"

            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            all_lines = list(reader)
            if not all_lines:
                raise HTTPException(status_code=400, detail="El archivo CSV está vacío.")

            header = [c.strip().lower() for c in all_lines[0]]
            code_idx, desc_idx, cif_idx = -1, -1, -1
            debe_idx, haber_idx, saldo_idx = -1, -1, -1

            for idx, h in enumerate(header):
                if any(k in h for k in ("cuenta", "codigo", "código", "subcuenta", "code")):
                    code_idx = idx
                elif any(k in h for k in ("descrip", "nombre", "concepto", "titulo", "name")):
                    desc_idx = idx
                elif any(k in h for k in ("cif", "nif", "dni", "fiscal")):
                    cif_idx = idx
                elif any(k in h for k in ("debe", "debit", "debe_acumulado", "debe_inicial")):
                    debe_idx = idx
                elif any(k in h for k in ("haber", "credit", "haber_acumulado", "haber_inicial")):
                    haber_idx = idx
                elif any(k in h for k in ("saldo", "balance")):
                    saldo_idx = idx

            start_line = 1 if (code_idx != -1 or desc_idx != -1) else 0
            if code_idx == -1:
                code_idx = 0
            if desc_idx == -1:
                desc_idx = 1 if len(all_lines[0]) > 1 else 0

            for line in all_lines[start_line:]:
                if not line or not line[code_idx].strip():
                    continue
                codigo_raw = line[code_idx].strip()
                desc_raw = line[desc_idx].strip() if desc_idx < len(line) else "Sin descripción"
                cif_raw = line[cif_idx].strip() if cif_idx != -1 and cif_idx < len(line) else ""

                d_val = parse_float_val(line[debe_idx]) if debe_idx != -1 and debe_idx < len(line) else 0.0
                h_val = parse_float_val(line[haber_idx]) if haber_idx != -1 and haber_idx < len(line) else 0.0
                if d_val == 0.0 and h_val == 0.0 and saldo_idx != -1 and saldo_idx < len(line):
                    s_val = parse_float_val(line[saldo_idx])
                    if s_val > 0:
                        d_val = s_val
                    elif s_val < 0:
                        h_val = abs(s_val)

                rows.append({
                    "codigo": codigo_raw,
                    "descripcion": desc_raw,
                    "cif": cif_raw,
                    "debe": d_val,
                    "haber": h_val,
                })
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Error leyendo archivo CSV: {str(e)}")

    # 2. Procesar e insertar masivamente
    created_count = 0
    updated_count = 0

    for item in rows:
        raw_code = item["codigo"].replace(".", "").replace(" ", "").strip()
        if len(raw_code) < expected_len:
            if len(raw_code) <= 4:
                continue
            pad_len = expected_len - len(raw_code)
            raw_code = raw_code[:3] + ("0" * pad_len) + raw_code[3:]
        elif len(raw_code) > expected_len:
            errors.append(f"Cuenta '{raw_code}' ignorada: supera los {expected_len} dígitos requeridos.")
            continue

        cif_norm = normalize_nif(item["cif"]) if item["cif"] else None
        tipo_deducido = deduce_account_type(raw_code)
        desc_clean = item["descripcion"][:250]
        debe_imp = item.get("debe", 0.0)
        haber_imp = item.get("haber", 0.0)

        # Comprobar si ya existe
        stmt = select(Account).where(Account.company_id == company_id, Account.codigo == raw_code)
        existing = (await db.execute(stmt)).scalars().first()

        if existing:
            existing.descripcion = desc_clean
            if cif_norm:
                existing.cif_asociado = cif_norm
            if debe_imp > 0:
                existing.debe_inicial = debe_imp
            if haber_imp > 0:
                existing.haber_inicial = haber_imp
            updated_count += 1
        else:
            acc = Account(
                company_id=company_id,
                codigo=raw_code,
                descripcion=desc_clean,
                tipo=tipo_deducido,
                cif_asociado=cif_norm,
                debe_inicial=debe_imp,
                haber_inicial=haber_imp,
            )
            db.add(acc)
            created_count += 1

        # Sincronizar proveedores si es 400.X o 410.X
        if tipo_deducido in ("PROVEEDOR", "ACREEDOR") and cif_norm:
            supp_stmt = select(Supplier).where(Supplier.company_id == company_id, Supplier.cif == cif_norm)
            supplier = (await db.execute(supp_stmt)).scalars().first()
            if not supplier:
                db.add(Supplier(
                    company_id=company_id,
                    cif=cif_norm,
                    nombre=desc_clean,
                    subcuenta_proveedor=raw_code,
                    subcuenta_gasto_defecto="629" + "0" * (expected_len - 4) + "1"
                ))
            else:
                supplier.subcuenta_proveedor = raw_code

    await db.commit()

    return AccountImportSummaryDTO(
        total_processed=len(rows),
        created=created_count,
        updated=updated_count,
        errors=errors[:15]
    )

@router.get("/{company_id}/chart-of-accounts/export")
async def export_chart_of_accounts(
    company_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Exporta el catálogo de cuentas con saldos contables acumulados en formato CSV."""
    c_res = await db.execute(select(Company).where(Company.id == company_id))
    company = c_res.scalars().first()
    if not company:
        raise HTTPException(status_code=404, detail="Empresa no encontrada.")

    res = await db.execute(
        select(Account).where(Account.company_id == company_id).order_by(Account.codigo.asc())
    )
    accounts = res.scalars().all()

    # Sumas acumuladas de apuntes
    entries_stmt = (
        select(
            AccountingEntryLine.subcuenta,
            func.coalesce(func.sum(AccountingEntryLine.debe), 0.0).label("sum_debe"),
            func.coalesce(func.sum(AccountingEntryLine.haber), 0.0).label("sum_haber")
        )
        .join(Invoice, Invoice.id == AccountingEntryLine.invoice_id)
        .where(Invoice.company_id == company_id)
        .group_by(AccountingEntryLine.subcuenta)
    )
    entries_res = await db.execute(entries_stmt)
    entries_map = {row.subcuenta: (float(row.sum_debe), float(row.sum_haber)) for row in entries_res.all()}

    output = io.StringIO()
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow([
        "Codigo",
        "Descripcion",
        "Tipo",
        "CIF_Asociado",
        "Debe_Acumulado",
        "Haber_Acumulado",
        "Saldo_Actual",
        "Tipo_Saldo"
    ])

    for a in accounts:
        e_debe, e_haber = entries_map.get(a.codigo, (0.0, 0.0))
        tot_debe = round(float(getattr(a, "debe_inicial", 0.0) or 0.0) + e_debe, 2)
        tot_haber = round(float(getattr(a, "haber_inicial", 0.0) or 0.0) + e_haber, 2)
        diff = round(tot_debe - tot_haber, 2)
        saldo = abs(diff)
        tipo_saldo = "DEUDOR" if diff > 0 else ("ACREEDOR" if diff < 0 else "CERO")

        writer.writerow([
            a.codigo,
            a.descripcion,
            a.tipo,
            a.cif_asociado or "",
            f"{tot_debe:.2f}",
            f"{tot_haber:.2f}",
            f"{saldo:.2f}",
            tipo_saldo
        ])

    return Response(
        content=output.getvalue().encode("utf-8"),
        media_type="text/csv",
        headers={
            "Content-Disposition": f"attachment; filename=plan_contable_{company.cif}.csv"
        }
    )
